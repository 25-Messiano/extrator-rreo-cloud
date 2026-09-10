from __future__ import annotations
import io, json, sys, time, tracemalloc
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from core.central_correcoes import PlanilhaEstadual, _read_state_sheet
from modules.mapeamento_nova_planilha import ABA_DESTINO

COUNTS={'AC':22,'AL':102,'AP':16,'AM':62,'BA':417,'CE':184,'DF':1,'ES':78,'GO':246,'MA':217,'MT':142,'MS':79,'MG':853,'PA':144,'PB':223,'PR':399,'PE':184,'PI':224,'RJ':92,'RN':167,'RS':497,'RO':52,'RR':15,'SC':295,'SP':645,'SE':75,'TO':139}

def make_xlsx(uf,n):
    wb=Workbook(); ws=wb.active; ws.title=ABA_DESTINO
    ws.cell(1,1).value='HOMOLOGACAO'
    prefix={'AC':'12','AL':'27','AP':'16','AM':'13','BA':'29','CE':'23','DF':'53','ES':'32','GO':'52','MA':'21','MT':'51','MS':'50','MG':'31','PA':'15','PB':'25','PR':'41','PE':'26','PI':'22','RJ':'33','RN':'24','RS':'43','RO':'11','RR':'14','SC':'42','SP':'35','SE':'28','TO':'17'}[uf]
    for i in range(n):
        r=i+3; ws.cell(r,3).value=(prefix+f'{i+1:05d}')[:7]; ws.cell(r,4).value=f'Municipio {i+1}/{uf}'
        for c in range(5,20): ws.cell(r,c).value=float(i+c)
    bio=io.BytesIO(); wb.save(bio); wb.close(); return bio.getvalue()

def main():
    tracemalloc.start(); started=time.perf_counter(); total=0; per={}
    for uf,n in COUNTS.items():
        payload=make_xlsx(uf,n)
        origin=PlanilhaEstadual(uf,'RREO',f'RREO_{uf}_2025_B6_RODADA_NOVA.xlsx',f'RODADAS/2025/RREO_{uf}.xlsx',datetime.now(timezone.utc))
        t=time.perf_counter(); rows=_read_state_sheet(payload,uf,'RREO',origin); dt=time.perf_counter()-t
        assert len(rows)==n,(uf,len(rows),n); total+=len(rows); per[uf]=round(dt,3)
    elapsed=time.perf_counter()-started; current,peak=tracemalloc.get_traced_memory(); tracemalloc.stop()
    result={'ufs':len(COUNTS),'municipios':total,'elapsed_s':round(elapsed,3),'peak_python_mb':round(peak/1024/1024,2),'slowest':sorted(per.items(),key=lambda x:x[1],reverse=True)[:5]}
    print(json.dumps(result,ensure_ascii=False))
    assert total==5570
    assert elapsed < 35.0
    return 0
if __name__=='__main__': raise SystemExit(main())
