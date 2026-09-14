from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from sqlalchemy import and_, or_, select

from database.db import session_scope
from models.entities import CodigoAPLB, Lancamento
from services.tenancy import get_active_tesouraria_id, tenant_where
from relatorios.movimento_financeiro import MESES, br_money, movimento_mes

STATUS_OFICIAIS = ("APROVADO", "CONCILIADO", "FECHADO")
MESES_CURTOS = {1:"JAN",2:"FEV",3:"MAR",4:"ABR",5:"MAI",6:"JUN",7:"JUL",8:"AGO",9:"SET",10:"OUT",11:"NOV",12:"DEZ"}
OFICIO = (216/25.4*72, 330/25.4*72)


def _norm_codigo(codigo: str | None) -> str:
    c = str(codigo or "").strip()
    return c.zfill(4) if c.isdigit() and len(c) <= 4 else c


def _rows_ano(ano: int) -> list[dict]:
    """Lê a base oficial por competência; data física é fallback para lançamentos sem competência."""
    cini, cfim = f"{ano:04d}-01", f"{ano:04d}-12"
    ini, fim = date(ano,1,1), date(ano,12,31)
    with session_scope() as s:
        q = (
            select(Lancamento, CodigoAPLB)
            .outerjoin(CodigoAPLB, CodigoAPLB.id == Lancamento.codigo_aplb_id)
            .where(
                Lancamento.status.in_(STATUS_OFICIAIS), tenant_where(Lancamento.tesouraria_id),
                or_(
                    and_(Lancamento.competencia >= cini, Lancamento.competencia <= cfim),
                    and_(Lancamento.competencia.is_(None), Lancamento.data_movimento >= ini, Lancamento.data_movimento <= fim),
                ),
            )
        )
        result = s.execute(q).all()
    out = []
    for l, c in result:
        try:
            mes = int(l.competencia[5:7]) if l.competencia else l.data_movimento.month
        except Exception:
            mes = l.data_movimento.month
        out.append({
            "id": l.id,
            "mes": mes,
            "codigo": _norm_codigo(c.codigo if c else ""),
            "descricao": c.descricao if c else (l.especificacao or ""),
            "natureza_padrao": c.natureza_padrao if c else None,
            "permite_entrada": bool(c.permite_entrada) if c else False,
            "permite_saida": bool(c.permite_saida) if c else False,
            "bc": l.origem_bc,
            "natureza": l.natureza,
            "valor": Decimal(l.valor),
        })
    return out


def _catalogo_completo(rows: list[dict]) -> list[dict]:
    """Catálogo ativo + códigos históricos usados, para não perder linha ao inativar cadastro."""
    with session_scope() as s:
        catalogo = [
            {"codigo": _norm_codigo(c.codigo), "descricao": c.descricao, "natureza": c.natureza_padrao or "", "permite_entrada": bool(c.permite_entrada), "permite_saida": bool(c.permite_saida), "ativo": bool(c.ativo)}
            for c in s.scalars(select(CodigoAPLB).where(CodigoAPLB.tesouraria_id==get_active_tesouraria_id()).order_by(CodigoAPLB.codigo)).all()
        ]
    vistos = {x["codigo"] for x in catalogo}
    for r in rows:
        if r["codigo"] and r["codigo"] not in vistos:
            catalogo.append({"codigo": r["codigo"], "descricao": r["descricao"], "natureza": r.get("natureza_padrao") or "", "permite_entrada": bool(r.get("permite_entrada")), "permite_saida": bool(r.get("permite_saida")), "ativo": False})
            vistos.add(r["codigo"])
    return catalogo


def resumo_codigo_ano(ano: int, bc: str | None = None) -> dict:
    """Resumo anual por código com Entradas/Saídas e reconciliação contra a base oficial."""
    rows = _rows_ano(ano)
    if bc:
        rows = [r for r in rows if r["bc"] == bc]
    catalogo = _catalogo_completo(rows)
    agg = defaultdict(lambda: defaultdict(lambda: {"ENTRADA": Decimal("0"), "SAIDA": Decimal("0")}))
    desc_usada = {}
    for r in rows:
        cod = r["codigo"] or "SEM CÓDIGO"
        agg[cod][r["mes"]][r["natureza"]] += r["valor"]
        desc_usada[cod] = r["descricao"]

    def montar(natureza: str) -> list[dict]:
        out = []
        for c in catalogo:
            # V22: a presença da linha no quadro é controlada por permissões
            # independentes definidas pelo administrador. Movimento histórico real
            # sempre permanece visível, mesmo se a permissão for alterada depois.
            tem_mov = any(agg[c["codigo"]][m][natureza] for m in range(1,13))
            permitido = bool(c.get("permite_entrada")) if natureza == "ENTRADA" else bool(c.get("permite_saida"))
            elegivel = permitido or tem_mov
            if not elegivel:
                continue
            item = {"codigo": c["codigo"], "descricao": c["descricao"] or desc_usada.get(c["codigo"], "")}
            item["meses"] = {m: agg[c["codigo"]][m][natureza] for m in range(1,13)}
            item["total"] = sum(item["meses"].values(), Decimal("0"))
            out.append(item)
        # SEM CÓDIGO, se houver, fica no final e é visível para auditoria.
        if "SEM CÓDIGO" in agg and any(agg["SEM CÓDIGO"][m][natureza] for m in range(1,13)):
            vals = {m: agg["SEM CÓDIGO"][m][natureza] for m in range(1,13)}
            out.append({"codigo":"SEM CÓDIGO","descricao":"LANÇAMENTOS SEM CÓDIGO","meses":vals,"total":sum(vals.values(),Decimal("0"))})
        return out

    entradas, saidas = montar("ENTRADA"), montar("SAIDA")
    total_ent = {m: sum((r["valor"] for r in rows if r["mes"] == m and r["natureza"] == "ENTRADA"), Decimal("0")) for m in range(1,13)}
    total_sai = {m: sum((r["valor"] for r in rows if r["mes"] == m and r["natureza"] == "SAIDA"), Decimal("0")) for m in range(1,13)}
    soma_ent = {m: sum((x["meses"][m] for x in entradas), Decimal("0")) for m in range(1,13)}
    soma_sai = {m: sum((x["meses"][m] for x in saidas), Decimal("0")) for m in range(1,13)}
    dif_ent = {m: total_ent[m] - soma_ent[m] for m in range(1,13)}
    dif_sai = {m: total_sai[m] - soma_sai[m] for m in range(1,13)}
    return {
        "ano": ano, "bc": bc, "entradas": entradas, "saidas": saidas,
        "totais_entrada": total_ent, "totais_saida": total_sai,
        "diferencas_entrada": dif_ent, "diferencas_saida": dif_sai,
        "ok": all(v == 0 for v in dif_ent.values()) and all(v == 0 for v in dif_sai.values()),
    }


def suprimento_caixa_ano(ano: int) -> dict:
    rows = [r for r in _rows_ano(ano) if r["codigo"] == "0801"]
    banco = {m: Decimal("0") for m in range(1,13)}
    caixa = {m: Decimal("0") for m in range(1,13)}
    for r in rows:
        if r["bc"] == "B" and r["natureza"] == "SAIDA":
            banco[r["mes"]] += r["valor"]
        elif r["bc"] == "C" and r["natureza"] == "ENTRADA":
            caixa[r["mes"]] += r["valor"]
    diff = {m: caixa[m] - banco[m] for m in range(1,13)}
    return {"ano": ano, "banco": banco, "caixa": caixa, "diferenca": diff, "ok": all(v == 0 for v in diff.values())}


def saldos_historicos(ano_ini: int, ano_fim: int) -> list[dict]:
    out = []
    for ano in range(ano_ini, ano_fim + 1):
        item = {"ano": ano, "B": {}, "C": {}, "T": {}}
        for m in range(1,13):
            mov = movimento_mes(ano, m)
            item["B"][m] = Decimal(mov["saldo_fim_banco"])
            item["C"][m] = Decimal(mov["saldo_fim_caixa"])
            item["T"][m] = Decimal(mov["saldo_fim_geral"])
        out.append(item)
    return out


def _money_excel(cell, value, negative=False):
    cell.value = float(-value if negative else value)
    cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00;R$ -'


def _style_table_excel(ws, start_row: int, end_row: int, end_col: int, title_row: int | None = None):
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.alignment = Alignment(vertical="center", wrap_text=True)
    if title_row:
        for c in ws[title_row]:
            c.font = Font(bold=True)


def gerar_excel_resumo(ano: int, bc: str | None, titulo: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    data = resumo_codigo_ano(ano, bc)
    for sem, meses in ((1, range(1,7)), (2, range(7,13))):
        ws = wb.create_sheet(f"{sem}o Semestre")
        ws.sheet_view.showGridLines = False
        r = 1
        for natureza, linhas, fill_title in (
            ("ENTRADAS", data["entradas"], "EAF4E3"),
            ("SAÍDAS", data["saidas"], "FCE4D6"),
        ):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
            ws.cell(r,1, f"{titulo}_{natureza} - {sem}º SEMESTRE DE {ano}")
            ws.cell(r,1).font = Font(bold=True, color="C00000")
            ws.cell(r,1).alignment = Alignment(horizontal="center")
            r += 1
            headers = ["CÓDIGO", f"REFERÊNCIA DAS {natureza}"] + [MESES[m] for m in meses] + ["TOTAL"]
            for j,h in enumerate(headers,1):
                ws.cell(r,j,h).font = Font(bold=True); ws.cell(r,j).alignment = Alignment(horizontal="center")
            head_row = r; r += 1
            for item in linhas:
                ws.cell(r,1,item["codigo"]); ws.cell(r,2,item["descricao"])
                for j,m in enumerate(meses,3): _money_excel(ws.cell(r,j), item["meses"][m], negative=(natureza=="SAÍDAS"))
                total = sum((item["meses"][m] for m in meses), Decimal("0"))
                _money_excel(ws.cell(r,9), total, negative=(natureza=="SAÍDAS")); r += 1
            ws.cell(r,2,"TOTAL..................................").font = Font(bold=True)
            totals = data["totais_entrada"] if natureza == "ENTRADAS" else data["totais_saida"]
            for j,m in enumerate(meses,3): _money_excel(ws.cell(r,j), totals[m], negative=(natureza=="SAÍDAS"))
            _money_excel(ws.cell(r,9), sum((totals[m] for m in meses),Decimal("0")), negative=(natureza=="SAÍDAS"))
            for c in ws[r]: c.font = Font(bold=True)
            end_row = r
            _style_table_excel(ws, head_row, end_row, 9)
            r += 3
        # resumo do semestre
        ws.cell(r,2,"CONFERÊNCIA").font = Font(bold=True)
        ws.cell(r,3,"ENTRADAS"); ws.cell(r,4,"SAÍDAS"); ws.cell(r,5,"RESULTADO")
        ent = sum((data["totais_entrada"][m] for m in meses),Decimal("0")); sai = sum((data["totais_saida"][m] for m in meses),Decimal("0"))
        _money_excel(ws.cell(r+1,3),ent); _money_excel(ws.cell(r+1,4),sai,negative=True); _money_excel(ws.cell(r+1,5),ent-sai)
        ws.cell(r+2,2,"Soma por códigos × Base oficial")
        ws.cell(r+2,3,"OK" if data["ok"] else "DIVERGÊNCIA")
        ws.cell(r+2,3).fill = PatternFill("solid", fgColor="C6EFCE" if data["ok"] else "FFC7CE")
        ws.column_dimensions['A'].width=12; ws.column_dimensions['B'].width=52
        for col in 'CDEFGHI': ws.column_dimensions[col].width=16
        ws.freeze_panes = "C3"
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize='14'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_margins.left=.2; ws.page_margins.right=.2; ws.page_margins.top=.3; ws.page_margins.bottom=.3
        ws.print_area=f"A1:I{ws.max_row}"
    out=BytesIO(); wb.save(out); return out.getvalue()


def _pdf_table_resumo(data, meses, natureza, titulo):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle
    small=ParagraphStyle('rr_small',fontName='Helvetica',fontSize=5.1,leading=5.8)
    hdr=['CÓDIGO',f'REFERÊNCIA DAS {natureza}']+[MESES[m].upper() for m in meses]+['TOTAL']
    rows=[hdr]
    lines=data['entradas'] if natureza=='ENTRADAS' else data['saidas']
    for item in lines:
        vals=[item['codigo'],Paragraph(item['descricao'],small)]
        for m in meses:
            v=item['meses'][m]
            vals.append(("-R$ " if natureza=='SAÍDAS' and v else "R$ ")+br_money(v))
        tot=sum((item['meses'][m] for m in meses),Decimal('0'))
        vals.append(("-R$ " if natureza=='SAÍDAS' and tot else "R$ ")+br_money(tot))
        rows.append(vals)
    totals=data['totais_entrada'] if natureza=='ENTRADAS' else data['totais_saida']
    vals=['','TOTAL..................................']
    for m in meses:
        v=totals[m]; vals.append(("-R$ " if natureza=='SAÍDAS' and v else "R$ ")+br_money(v))
    tot=sum((totals[m] for m in meses),Decimal('0')); vals.append(("-R$ " if natureza=='SAÍDAS' and tot else "R$ ")+br_money(tot)); rows.append(vals)
    t=Table(rows,repeatRows=1,colWidths=[40,270]+[65]*7)
    style=[('FONT',(0,0),(-1,-1),'Helvetica',5.1),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),0.35,colors.black),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(2,1),(-1,-1),'RIGHT'),('FONTNAME',(0,-1),(-1,-1),'Helvetica-Bold')]
    if natureza=='SAÍDAS': style += [('TEXTCOLOR',(2,1),(-1,-1),colors.red)]
    else: style += [('TEXTCOLOR',(2,1),(-1,-1),colors.HexColor('#1599C7'))]
    t.setStyle(TableStyle(style)); return t


def gerar_pdf_resumo(ano: int, bc: str | None, titulo: str) -> bytes:
    from reportlab.lib.pagesizes import landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib import colors
    data=resumo_codigo_ano(ano,bc); buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(OFICIO),leftMargin=14,rightMargin=14,topMargin=14,bottomMargin=14)
    styles=getSampleStyleSheet(); h=ParagraphStyle('h',parent=styles['Heading2'],fontSize=9,leading=10,alignment=1,textColor=colors.red)
    note=ParagraphStyle('n',parent=styles['BodyText'],fontSize=7,leading=8)
    story=[]
    for idx,(sem,meses) in enumerate(((1,list(range(1,7))),(2,list(range(7,13))))):
        if idx: story.append(PageBreak())
        story.append(Paragraph(f'{titulo} - {sem}º SEMESTRE DE {ano}',h)); story.append(Spacer(1,4))
        story.append(_pdf_table_resumo(data,meses,'ENTRADAS',titulo)); story.append(Spacer(1,8))
        story.append(_pdf_table_resumo(data,meses,'SAÍDAS',titulo)); story.append(Spacer(1,8))
        ent=sum((data['totais_entrada'][m] for m in meses),Decimal('0')); sai=sum((data['totais_saida'][m] for m in meses),Decimal('0'))
        story.append(Paragraph(f'ENTRADAS: R$ {br_money(ent)} &nbsp;&nbsp;&nbsp; SAÍDAS: -R$ {br_money(sai)} &nbsp;&nbsp;&nbsp; RESULTADO: R$ {br_money(ent-sai)}',note))
        story.append(Paragraph('CONFERÊNCIA SOMA DOS CÓDIGOS × BASE OFICIAL: '+('OK' if data['ok'] else 'DIVERGÊNCIA'),note))
    doc.build(story); return buf.getvalue()


def gerar_excel_suprimento(ano:int)->bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    d=suprimento_caixa_ano(ano); wb=Workbook(); wb.remove(wb.active)
    for sem,meses in ((1,range(1,7)),(2,range(7,13))):
        ws=wb.create_sheet(f'{sem}o Semestre'); ws.sheet_view.showGridLines=False
        headers=[f'ANO {ano}','COD']+[MESES[m] for m in meses]+['TOTAL'];
        for j,h in enumerate(headers,1): ws.cell(1,j,h).font=Font(bold=True)
        ws.cell(2,1,'BANCO......'); ws.cell(2,2,'0801'); ws.cell(3,1,'CAIXA......'); ws.cell(3,2,'0801'); ws.cell(4,1,'DIFERENÇA')
        for j,m in enumerate(meses,3):
            _money_excel(ws.cell(2,j),d['banco'][m],negative=True); _money_excel(ws.cell(3,j),d['caixa'][m]); _money_excel(ws.cell(4,j),d['diferenca'][m])
        _money_excel(ws.cell(2,9),sum((d['banco'][m] for m in meses),Decimal('0')),negative=True)
        _money_excel(ws.cell(3,9),sum((d['caixa'][m] for m in meses),Decimal('0')))
        _money_excel(ws.cell(4,9),sum((d['diferenca'][m] for m in meses),Decimal('0')))
        for c in ws[2][2:]: c.font=Font(color='FF0000',bold=True)
        for c in ws[3][2:]: c.font=Font(color='1599C7',bold=True)
        ws['A6']='CONFERÊNCIA'; ws['B6']='OK' if all(d['diferenca'][m]==0 for m in meses) else 'DIVERGÊNCIA'; ws['B6'].fill=PatternFill('solid',fgColor='C6EFCE' if ws['B6'].value=='OK' else 'FFC7CE')
        _style_table_excel(ws,1,4,9); ws.column_dimensions['A'].width=18;ws.column_dimensions['B'].width=10
        for col in 'CDEFGHI': ws.column_dimensions[col].width=16
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize='14';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO();wb.save(out);return out.getvalue()


def gerar_pdf_suprimento(ano:int)->bytes:
    from reportlab.lib.pagesizes import landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    d=suprimento_caixa_ano(ano);buf=BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(OFICIO),leftMargin=18,rightMargin=18,topMargin=18,bottomMargin=18);styles=getSampleStyleSheet();story=[]
    for idx,(sem,meses) in enumerate(((1,list(range(1,7))),(2,list(range(7,13))))):
        if idx: story.append(PageBreak())
        story.append(Paragraph(f'SUPRIMENTO DE CAIXA - {sem}º SEMESTRE DE {ano}',styles['Heading2'])); story.append(Spacer(1,6))
        rows=[[f'ANO {ano}','COD']+[MESES[m] for m in meses]+['TOTAL']]
        b=[sum((d['banco'][m] for m in meses),Decimal('0'))];c=[sum((d['caixa'][m] for m in meses),Decimal('0'))];dif=[sum((d['diferenca'][m] for m in meses),Decimal('0'))]
        rows += [['BANCO......','0801']+['-R$ '+br_money(d['banco'][m]) if d['banco'][m] else 'R$ -' for m in meses]+['-R$ '+br_money(b[0]) if b[0] else 'R$ -'],['CAIXA......','0801']+['R$ '+br_money(d['caixa'][m]) if d['caixa'][m] else 'R$ -' for m in meses]+['R$ '+br_money(c[0]) if c[0] else 'R$ -'],['DIFERENÇA','']+['R$ '+br_money(d['diferenca'][m]) for m in meses]+['R$ '+br_money(dif[0])]]
        t=Table(rows,colWidths=[80,50]+[90]*7);t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.5,colors.black),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('ALIGN',(2,1),(-1,-1),'RIGHT'),('TEXTCOLOR',(2,1),(-1,1),colors.red),('TEXTCOLOR',(2,2),(-1,2),colors.HexColor('#1599C7')),('FONT',(0,0),(-1,-1),'Helvetica',8)]));story.append(t);story.append(Spacer(1,8));story.append(Paragraph('Conferência Banco × Caixa: '+('OK' if all(d['diferenca'][m]==0 for m in meses) else 'DIVERGÊNCIA'),styles['BodyText']))
    doc.build(story);return buf.getvalue()


def gerar_excel_saldos(ano_ini:int,ano_fim:int)->bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows=saldos_historicos(ano_ini,ano_fim);wb=Workbook();ws=wb.active;ws.title='Saldos Históricos';ws.sheet_view.showGridLines=False;r=1
    for key,label in [('B','SALDOS BANCO'),('C','SALDOS CAIXA'),('T','SALDOS: BANCO + CAIXA')]:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=13);ws.cell(r,1,f'{label} - ANOS {ano_ini} A {ano_fim}').font=Font(bold=True);r+=1
        for j,h in enumerate(['ANO']+[MESES[m] for m in range(1,13)],1):ws.cell(r,j,h).font=Font(bold=True)
        start=r;r+=1
        for item in rows:
            ws.cell(r,1,item['ano'])
            for m in range(1,13):_money_excel(ws.cell(r,m+1),item[key][m])
            if item['ano']%2==0:
                for c in ws[r]:c.fill=PatternFill('solid',fgColor='FFF200')
            r+=1
        _style_table_excel(ws,start,r-1,13);r+=2
    ws.column_dimensions['A'].width=10
    from openpyxl.utils import get_column_letter
    for col in range(2,14): ws.column_dimensions[get_column_letter(col)].width=15
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize='14';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO();wb.save(out);return out.getvalue()


def gerar_pdf_saldos(ano_ini:int,ano_fim:int)->bytes:
    from reportlab.lib.pagesizes import landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    rows=saldos_historicos(ano_ini,ano_fim);buf=BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(OFICIO),leftMargin=14,rightMargin=14,topMargin=14,bottomMargin=14);styles=getSampleStyleSheet();story=[]
    for key,label in [('B','SALDOS BANCO'),('C','SALDOS CAIXA'),('T','SALDOS: BANCO + CAIXA')]:
        story.append(Paragraph(f'{label} - ANOS {ano_ini} A {ano_fim}',styles['Heading3']))
        data=[['ANO']+[MESES_CURTOS[m] for m in range(1,13)]]+[[x['ano']]+['R$ '+br_money(x[key][m]) for m in range(1,13)] for x in rows]
        t=Table(data,repeatRows=1,colWidths=[42]+[61]*12);style=[('GRID',(0,0),(-1,-1),0.4,colors.black),('FONT',(0,0),(-1,-1),'Helvetica',5.8),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('ALIGN',(1,1),(-1,-1),'RIGHT')]
        for i,x in enumerate(rows,1):
            if x['ano']%2==0: style.append(('BACKGROUND',(0,i),(-1,i),colors.yellow))
        t.setStyle(TableStyle(style));story+=[t,Spacer(1,10)]
    doc.build(story);return buf.getvalue()
