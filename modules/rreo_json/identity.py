from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl.worksheet.worksheet import Worksheet

from modules.rreo_safe import canonical_municipality_name, extract_municipality_header_candidates, load_municipality_aliases


def normalize_name(value: Any) -> str:
    raw = str(value or "").strip()
    raw = re.sub(r"\b([Dd])['’`´]\s*([A-Za-zÀ-ÿ])", r"\1\2", raw)
    raw = unicodedata.normalize("NFKD", raw)
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    raw = raw.upper().replace("–", "-").replace("—", "-")
    raw = re.sub(r"[^A-Z0-9]+", " ", raw)
    return re.sub(r"\s+", " ", raw).strip()


def external_name_from_filename(filename: str, uf: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"(?i)^RREO[_\s-]*MUNICIPAL[_\s-]*\d{4}A?[_\s-]*", "", stem).strip(" _-")
    stem = re.sub(rf"(?i)\s*[-_ ]\s*{re.escape(uf)}\s*$", "", stem).strip(" _-")
    return stem


def _tokens(value: str) -> set[str]:
    return {x for x in normalize_name(value).split() if x}


def similarity(a: str, b: str) -> float:
    na, nb = normalize_name(a), normalize_name(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    ta, tb = _tokens(na), _tokens(nb)
    token_score = len(ta & tb) / max(1, len(ta | tb))
    seq = SequenceMatcher(None, na, nb).ratio()
    return round(max(seq, (0.55 * seq + 0.45 * token_score)), 6)


def _split_ente(value: Any) -> tuple[str, str]:
    raw = str(value or "").strip()
    m = re.match(r"^(.*?)\s*/\s*([A-Za-z]{2})\s*$", raw)
    if m:
        return m.group(1).strip(), m.group(2).upper()
    return raw, ""


@dataclass(frozen=True)
class Candidate:
    row: int
    ibge: str
    ente: str
    municipio: str
    uf: str


def worksheet_candidates(ws: Worksheet, uf: str, *, ibge_col: int = 3, ente_col: int = 4, first_row: int = 3) -> list[Candidate]:
    uf = str(uf or "").upper().strip()
    out: list[Candidate] = []
    for row in range(first_row, ws.max_row + 1):
        ente = str(ws.cell(row, ente_col).value or "").strip()
        municipio, row_uf = _split_ente(ente)
        if row_uf and row_uf != uf:
            continue
        ibge = re.sub(r"\D", "", str(ws.cell(row, ibge_col).value or ""))
        if not municipio:
            continue
        out.append(Candidate(row=row, ibge=ibge, ente=ente, municipio=municipio, uf=row_uf or uf))
    return out


def _alias_canonical(name: str, uf: str) -> str:
    key = canonical_municipality_name(name)
    official = load_municipality_aliases().get((uf.upper(), key))
    return official or key


def _score_candidate(external: str, internal: str, candidate: Candidate, uf: str) -> tuple[float, float, float]:
    target = candidate.municipio
    ext_canon = _alias_canonical(external, uf)
    int_canon = _alias_canonical(internal, uf) if internal else ""
    tgt_canon = _alias_canonical(target, uf)
    s_ext = 1.0 if ext_canon and ext_canon == tgt_canon else similarity(external, target)
    s_int = 1.0 if int_canon and int_canon == tgt_canon else similarity(internal, target) if internal else 0.0
    if external and internal:
        combined = (0.55 * s_ext) + (0.45 * s_int)
        # convergence bonus: both independently point strongly to same ente
        if s_ext >= 0.90 and s_int >= 0.90:
            combined = min(1.0, combined + 0.03)
    elif external:
        combined = s_ext
    else:
        combined = s_int
    return round(combined, 6), round(s_ext, 6), round(s_int, 6)


def identify_municipality(
    *,
    filename: str,
    internal_text: str,
    ws: Worksheet,
    uf: str,
    similarity_min: float = 0.82,
    ambiguity_margin: float = 0.03,
) -> dict[str, Any]:
    uf = str(uf or "").upper().strip()
    external = external_name_from_filename(filename, uf)
    internal_candidates = extract_municipality_header_candidates(internal_text, uf)
    internal = internal_candidates[0] if internal_candidates else ""
    candidates = worksheet_candidates(ws, uf)
    ranked: list[dict[str, Any]] = []
    for c in candidates:
        score, s_ext, s_int = _score_candidate(external, internal, c, uf)
        ranked.append({
            "row": c.row,
            "ibge": c.ibge,
            "ente_planilha": c.ente,
            "municipio_planilha": c.municipio,
            "uf": c.uf,
            "score": score,
            "score_externo": s_ext,
            "score_interno": s_int,
        })
    ranked.sort(key=lambda x: (-x["score"], -x["score_externo"], -x["score_interno"], x["row"]))
    best = ranked[0] if ranked else None
    second = ranked[1] if len(ranked) > 1 else None
    if not best:
        status = "SEM_CANDIDATO_UF"
    elif best["score"] < similarity_min:
        status = "PENDENTE_CONFERENCIA"
    elif second and (best["score"] - second["score"]) < ambiguity_margin:
        status = "AMBIGUO"
    else:
        exact_ext = normalize_name(external) == normalize_name(best["municipio_planilha"])
        exact_int = bool(internal) and normalize_name(internal) == normalize_name(best["municipio_planilha"])
        if exact_ext or exact_int:
            status = "EXATO"
        elif best["score"] >= 0.94:
            status = "ALTA_CONFIANCA"
        else:
            status = "TOLERANCIA_VALIDADA"
    return {
        "tipo": "IDENTIDADE_MUNICIPIO",
        "uf": uf,
        "nome_externo": external,
        "nome_externo_normalizado": normalize_name(external),
        "nome_interno": internal,
        "nome_interno_normalizado": normalize_name(internal),
        "candidatos_internos": internal_candidates,
        "status": status,
        "identificado": status in {"EXATO", "ALTA_CONFIANCA", "TOLERANCIA_VALIDADA"},
        "melhor_candidato": best,
        "segundo_candidato": second,
        "top5": ranked[:5],
        "regra": "NOME_EXTERNO+NOME_INTERNO+ENTE_PLANILHA+UF; IBGE SOMENTE APOS HARMONIZACAO",
        "limiar_similaridade": similarity_min,
        "margem_ambiguidade": ambiguity_margin,
    }
