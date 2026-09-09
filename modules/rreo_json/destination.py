from __future__ import annotations

import re
from typing import Any, Mapping

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from modules.rreo_destinos import CAMPOS, POR_CODIGO, validar_destinos_unicos


def build_destination_record(ws: Worksheet, identity: Mapping[str, Any], validation: Mapping[str, Any], values: Mapping[str, Any]) -> dict[str, Any]:
    if validation.get("status") != "VALIDADO_SAFE_JSON":
        raise RuntimeError("Destino bloqueado: JSON de validacao nao esta VALIDADO_SAFE_JSON")
    candidate = identity.get("melhor_candidato") or {}
    row = int(candidate.get("row") or 0)
    ibge = re.sub(r"\D", "", str(candidate.get("ibge") or ""))
    ente = str(candidate.get("ente_planilha") or "")
    if row < 1 or not ente:
        raise RuntimeError("Identidade sem linha/ente confirmado")

    header_map = validar_destinos_unicos(ws, CAMPOS)
    destinations: dict[str, Any] = {}
    for code, value in values.items():
        field = POR_CODIGO.get(str(code))
        if field is None:
            continue
        col = header_map.get(str(code))
        if not field.gravar:
            destinations[str(code)] = {"modo": "VALIDACAO_SOMENTE", "valor": value}
            continue
        if value is None:
            destinations[str(code)] = {"modo": "SEM_VALOR", "valor": None}
            continue
        if col is None:
            raise RuntimeError(f"Destino nao resolvido para {code}")
        destinations[str(code)] = {
            "modo": "GRAVAR",
            "valor": round(float(value), 2),
            "linha": row,
            "coluna": col,
            "letra": get_column_letter(col),
            "celula": f"{get_column_letter(col)}{row}",
            "cabecalho": ws.cell(1, col).value,
        }
    return {
        "tipo": "DESTINO_RREO",
        "aba": ws.title,
        "linha": row,
        "ibge_registrado": ibge,
        "ente_planilha": ente,
        "destinos": destinations,
        "regra": "DESTINO RESOLVIDO PELO CABECALHO; LINHA VEM DA HARMONIZACAO NOMINAL; IBGE E REGISTRO DE CONFIRMACAO",
    }


def apply_destination_record(ws: Worksheet, destination: Mapping[str, Any], validation: Mapping[str, Any]) -> list[str]:
    if validation.get("status") != "VALIDADO_SAFE_JSON":
        raise RuntimeError("Gravacao bloqueada: validacao insuficiente")
    expected_sheet = destination.get("aba")
    if ws.title != expected_sheet:
        raise RuntimeError(f"Aba divergente: {ws.title!r} != {expected_sheet!r}")
    written: list[str] = []
    for code, item in (destination.get("destinos") or {}).items():
        if item.get("modo") != "GRAVAR":
            continue
        row, col = int(item["linha"]), int(item["coluna"])
        if str(ws.cell(1, col).value or "") != str(item.get("cabecalho") or ""):
            raise RuntimeError(f"Cabecalho mudou para {code}; gravacao abortada")
        cell = ws.cell(row, col)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise RuntimeError(f"Formula protegida em {cell.coordinate}")
        cell.value = float(item["valor"])
        cell.number_format = '#,##0.00;[Red](#,##0.00);-'
        written.append(cell.coordinate)
    return written
