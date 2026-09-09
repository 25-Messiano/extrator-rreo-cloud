from __future__ import annotations
import json
from pathlib import Path
from openpyxl import load_workbook
from .coluna_a_linhas import validar as validar_a
from .coluna_b_municipios import validar as validar_b
from .coluna_c_ibge import validar as validar_c, normalizar as ibge_normalizar
from .coluna_d_ente_federado import validar as validar_d


def _eh_estado(ente: str, ibge: str) -> bool:
    return str(ente or "").strip().lower().startswith("estado d") and len(ibge) == 2


def auditar(planilha: str | Path, saida_json: str | Path | None = None) -> dict:
    wb = load_workbook(planilha, data_only=False, read_only=True)
    ws = wb[wb.sheetnames[0]]
    erros = []
    municipios = []
    estados = []
    esperado_a = 1
    esperado_b = None
    uf_atual = ""

    for r, values in enumerate(ws.iter_rows(min_row=3, min_col=1, max_col=4, values_only=True), start=3):
        a, b, c, d = values
        if all(v in (None, "") for v in (a, b, c, d)):
            continue
        ibge = ibge_normalizar(c)
        vd = validar_d(d)
        estado = _eh_estado(d, ibge)
        municipio = len(ibge) == 7 and bool(vd["uf"]) and not estado

        va = validar_a(a, esperado_a)
        if not va["ok"]:
            erros.append({"linha_excel": r, "tipo": "COLUNA_A", **va})
        esperado_a += 1

        if estado:
            uf_atual = vd["uf"]
            esperado_b = 1
            vb = validar_b(b, None, True)
            estados.append({"linha_excel": r, "ibge_uf": ibge, "ente": d, "uf": uf_atual})
        elif municipio:
            vb = validar_b(b, esperado_b, False)
            esperado_b = (esperado_b or 1) + 1
            vc = validar_c(c, True)
            if vd["uf"] and uf_atual and vd["uf"] != uf_atual:
                erros.append({"linha_excel": r, "tipo": "UF_FORA_DO_BLOCO", "uf_linha": vd["uf"], "uf_bloco": uf_atual})
            municipios.append({
                "linha_excel": r,
                "sequencial_geral": a,
                "sequencial_uf": b,
                "ibge": vc["ibge"],
                "ente_federado": d,
                "nome": vd["nome"],
                "nome_normalizado": vd["nome_normalizado"],
                "uf": vd["uf"],
            })
            if not vc["ok"]:
                erros.append({"linha_excel": r, "tipo": "COLUNA_C", **vc})
        else:
            vb = {"ok": True, "coluna": "B", "papel": "NAO_APLICAVEL", "valor": b, "esperado": None}

        if not vb["ok"]:
            erros.append({"linha_excel": r, "tipo": "COLUNA_B", **vb})
        if not vd["ok"]:
            erros.append({"linha_excel": r, "tipo": "COLUNA_D", **vd})

    resultado = {
        "versao": "V1.3.9",
        "arquivo": Path(planilha).name,
        "aba": ws.title,
        "max_row_excel": ws.max_row,
        "max_column_excel": ws.max_column,
        "quantidade_linhas_cadastrais_coluna_a": esperado_a - 1,
        "quantidade_municipios_coluna_b": len(municipios),
        "quantidade_estados_detectados": len(estados),
        "estrutura_abcd_valida": len(erros) == 0,
        "erros": erros,
        "estados": estados,
        "municipios": municipios,
    }
    wb.close()
    if saida_json:
        Path(saida_json).write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding="utf-8")
    return resultado


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("planilha")
    p.add_argument("--json", default="estrutura_abcd.json")
    args = p.parse_args()
    res = auditar(args.planilha, args.json)
    print(json.dumps({k: v for k, v in res.items() if k not in {"municipios", "estados", "erros"}}, ensure_ascii=False, indent=2))
    print(f"erros={len(res['erros'])}")
