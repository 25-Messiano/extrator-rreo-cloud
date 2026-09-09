from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

CONFIG_PADRAO = Path(__file__).resolve().parents[2] / "config" / "mapeamento_rreo_excel.json"


@dataclass(frozen=True)
class Destination:
    codigo: str
    modo: str
    coluna: str | None
    linha: int | None
    celula: str | None
    ibge: str | None
    ente_federado: str | None


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper().replace("–", "-").replace("—", "-")
    text = re.sub(r"[^A-Z0-9.]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def exact_code_in_header(header: Any, code: str) -> bool:
    raw = str(header or "")
    pattern = r"(?<![0-9.])" + re.escape(code) + r"(?![0-9.])"
    return bool(re.search(pattern, raw))


def load_config(path: str | Path = CONFIG_PADRAO) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as fh:
        cfg = json.load(fh)
    validate_config(cfg)
    return cfg


def validate_config(cfg: Mapping[str, Any]) -> None:
    structure = cfg.get("estrutura_planilha", {})
    if tuple(structure.keys()) != ("A", "B", "C", "D"):
        raise ValueError("Estrutura obrigatoria deve ser A, B, C, D nesta ordem")
    cols = cfg.get("colunas_pdf", {})
    if cols.get("a", {}).get("coletar") is not False or cols.get("a", {}).get("gravar_excel") is not False:
        raise ValueError("PREVISAO ATUALIZADA (a) deve permanecer bloqueada")
    if cols.get("b", {}).get("coletar") is not True or cols.get("b", {}).get("gravar_excel") is not True:
        raise ValueError("RECEITAS REALIZADAS (b) deve permanecer autorizada")

    used: dict[str, str] = {}
    for code, rule in cfg.get("codigos", {}).items():
        mode = rule.get("modo")
        dest = rule.get("destino")
        if mode not in {"gravar", "validar"}:
            raise ValueError(f"Modo invalido em {code}: {mode}")
        if mode == "validar":
            if dest is not None:
                raise ValueError(f"{code} e validador e nao pode ter destino")
            continue
        if not dest or not re.fullmatch(r"[A-Z]+", str(dest)):
            raise ValueError(f"Destino invalido em {code}: {dest}")
        if dest in used:
            raise ValueError(f"Colisao de destino {dest}: {used[dest]} e {code}")
        used[dest] = code


def select_sheet(workbook, cfg: Mapping[str, Any]) -> Worksheet:
    # O modelo atual usa uma unica aba com nome historico; o nome da aba nao e autoridade.
    return workbook.active


def _header_map(ws: Worksheet, cfg: Mapping[str, Any]) -> dict[str, str | None]:
    header_row = int(cfg["planilha"].get("linha_cabecalho", 1))
    actual: dict[str, str | None] = {}
    occupied: dict[int, str] = {}
    for code, rule in cfg["codigos"].items():
        if rule["modo"] == "validar":
            actual[code] = None
            continue
        matches: list[int] = []
        for col in range(1, ws.max_column + 1):
            if exact_code_in_header(ws.cell(header_row, col).value, code):
                matches.append(col)
        if len(matches) != 1:
            raise RuntimeError(f"Cabecalho {code}: esperado 1 destino, encontrados {len(matches)}: {matches}")
        col = matches[0]
        expected = str(rule["destino"])
        real = get_column_letter(col)
        if real != expected:
            raise RuntimeError(f"Cabecalho {code} mudou: configurado {expected}, encontrado {real}")
        if col in occupied:
            raise RuntimeError(f"Colisao real de cabecalho: {code} e {occupied[col]} em {real}")
        occupied[col] = code
        actual[code] = real
    return actual


def validate_headers(ws: Worksheet, cfg: Mapping[str, Any] | None = None) -> dict[str, str | None]:
    cfg = dict(cfg or load_config())
    expected_struct = {"A": "Nº", "B": "Qut. Municípios", "C": "Código IBGE", "D": "Ente Federado"}
    for letter, expected in expected_struct.items():
        value = ws[f"{letter}1"].value
        if normalize(value) != normalize(expected):
            raise RuntimeError(f"Cabecalho estrutural {letter} invalido: {value!r}; esperado {expected!r}")
    if ws.max_column < int(cfg["planilha"].get("total_colunas_esperado", 27)):
        raise RuntimeError(f"Planilha tem somente {ws.max_column} colunas; esperado >= 27")
    return _header_map(ws, cfg)


def validate_identity_row(ws: Worksheet, identity: Mapping[str, Any]) -> dict[str, Any]:
    if not identity.get("identificado"):
        raise RuntimeError(f"Identidade nao liberada: {identity.get('status')}")
    best = identity.get("melhor_candidato") or {}
    row = int(best.get("row") or 0)
    if row < 3:
        raise RuntimeError(f"Linha de municipio invalida: {row}")
    ente = str(ws[f"D{row}"].value or "").strip()
    ibge = re.sub(r"\D", "", str(ws[f"C{row}"].value or ""))
    if ente != str(best.get("ente_planilha") or "").strip():
        raise RuntimeError(f"Ente mudou entre identidade e destino: {best.get('ente_planilha')!r} != {ente!r}")
    if len(ibge) != 7:
        raise RuntimeError(f"IBGE estrutural invalido em C{row}: {ibge!r}")
    if best.get("ibge") and re.sub(r"\D", "", str(best.get("ibge"))) != ibge:
        raise RuntimeError(f"IBGE mudou entre identidade e destino em linha {row}")
    # A e B devem existir para linhas municipais.
    a = ws[f"A{row}"].value
    b = ws[f"B{row}"].value
    if not isinstance(a, (int, float)) or not isinstance(b, (int, float)):
        raise RuntimeError(f"Estrutura A/B invalida na linha {row}: A={a!r}, B={b!r}")
    return {"linha": row, "ibge": ibge, "ente_federado": ente, "A": int(a), "B": int(b)}


def resolve_destination(ws: Worksheet, identity: Mapping[str, Any], code: str, cfg: Mapping[str, Any] | None = None) -> Destination:
    cfg = dict(cfg or load_config())
    validate_headers(ws, cfg)
    code = str(code).strip()
    if code not in cfg["codigos"]:
        raise KeyError(f"Codigo RREO nao configurado: {code}")
    rule = cfg["codigos"][code]
    if rule["modo"] == "validar":
        return Destination(code, "validar", None, None, None, None, None)
    row_info = validate_identity_row(ws, identity)
    col = str(rule["destino"])
    if col in set(cfg.get("colunas_proibidas_rreo", [])):
        raise RuntimeError(f"Destino proibido para RREO: {col}")
    return Destination(
        codigo=code,
        modo="gravar",
        coluna=col,
        linha=row_info["linha"],
        celula=f"{col}{row_info['linha']}",
        ibge=row_info["ibge"],
        ente_federado=row_info["ente_federado"],
    )


def build_destination_json(ws: Worksheet, identity: Mapping[str, Any], validated_values: Mapping[str, Any], cfg: Mapping[str, Any] | None = None) -> dict[str, Any]:
    cfg = dict(cfg or load_config())
    validate_headers(ws, cfg)
    destinations: dict[str, Any] = {}
    for code, value in validated_values.items():
        if code not in cfg["codigos"]:
            continue
        dest = resolve_destination(ws, identity, code, cfg)
        output_mode = "VALIDACAO_SOMENTE" if dest.modo == "validar" else "GRAVAR"
        destinations[code] = {
            "modo": output_mode,
            "valor": value,
            "coluna": dest.coluna,
            "linha": dest.linha,
            "celula": dest.celula,
            "ibge": dest.ibge,
            "ente_federado": dest.ente_federado,
            "autorizado_gravar": dest.modo == "gravar" and dest.celula is not None,
        }
    return {
        "tipo": "DESTINO_RREO_INTEGRADO",
        "versao": cfg["versao"],
        "status": "PRONTO_PARA_IMPLANTAR",
        "destinos": destinations,
        "regra": "IDENTIDADE -> A-D -> CABECALHO REAL -> CODIGO -> CELULA; EXTRAÇÃO NAO ESCREVE EXCEL",
    }


def apply_destination_json(workbook_path: str | Path, destination_record: Mapping[str, Any], output_path: str | Path | None = None) -> Path:
    src = Path(workbook_path)
    out = Path(output_path) if output_path else src
    wb = load_workbook(src, data_only=False)
    try:
        ws = wb.active
        cfg = load_config()
        validate_headers(ws, cfg)
        for code, item in destination_record.get("destinos", {}).items():
            if not item.get("autorizado_gravar"):
                continue
            cell = str(item["celula"])
            expected_col = str(cfg["codigos"][code]["destino"])
            if re.match(r"^[A-Z]+", cell).group(0) != expected_col:
                raise RuntimeError(f"Celula adulterada para {code}: {cell}")
            ws[cell] = float(item["valor"])
            ws[cell].number_format = '#,##0.00;[Red](#,##0.00);-'
        wb.save(out)
    finally:
        wb.close()
    return out
