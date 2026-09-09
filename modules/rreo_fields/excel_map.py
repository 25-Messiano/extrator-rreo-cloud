from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Mapping

from openpyxl.worksheet.worksheet import Worksheet

from modules.mapeamento_nova_planilha import ABA_DESTINO

# Identidade fixa da planilha oficial.
IBGE_COLUMN = 3       # C
MUNICIPIO_COLUMN = 4  # D
FIRST_MUNICIPAL_ROW = 3

# Somente colunas realmente existentes na planilha-base oficial.
# 2.1 e 6.2 sao totais de VALIDACAO, nao possuem destino oficial e nao sao inventados.
CODE_TO_EXCEL: dict[str, tuple[int, str] | None] = {
    "1.1": (16, "P"),
    "1.2": (18, "R"),
    "1.3": (19, "S"),
    "1.4": (17, "Q"),
    "2.1": None,
    "2.1.1": (5, "E"),
    "2.1.2": (6, "F"),
    "2.2": (11, "K"),
    "2.3": (8, "H"),
    "2.4": (10, "J"),
    "2.5": (12, "L"),
    "2.6": (20, "T"),
    "6.1.1": (14, "N"),
    "6.2": None,
    "6.2.1": (15, "O"),
}


def normalize_ibge(value: Any) -> str:
    digits = "".join(ch for ch in str(value or "") if ch.isdigit())
    return digits.zfill(7) if digits else ""


def build_ibge_row_index(ws: Worksheet) -> dict[str, int]:
    if ws.title != ABA_DESTINO:
        raise ValueError(f"Aba incorreta: {ws.title!r}; esperada: {ABA_DESTINO!r}")
    out: dict[str, int] = {}
    for row_number, cells in enumerate(
        ws.iter_rows(
            min_row=FIRST_MUNICIPAL_ROW,
            max_row=ws.max_row,
            min_col=IBGE_COLUMN,
            max_col=IBGE_COLUMN,
            values_only=True,
        ),
        start=FIRST_MUNICIPAL_ROW,
    ):
        ibge = normalize_ibge(cells[0])
        if len(ibge) == 7:
            out[ibge] = row_number
    return out


def destination_for(ibge: str, code: str, index: Mapping[str, int]) -> tuple[int, int, str] | None:
    dest = CODE_TO_EXCEL.get(code)
    if dest is None:
        return None
    row = index.get(normalize_ibge(ibge))
    if row is None:
        raise KeyError(f"IBGE nao encontrado na planilha: {ibge}")
    return row, dest[0], dest[1]


def write_confirmed(ws: Worksheet, ibge: str, values: Mapping[str, float | None]) -> dict[str, object]:
    index = build_ibge_row_index(ws)
    written: list[str] = []
    validation_only: list[str] = []
    for code, value in values.items():
        if value is None:
            continue
        destination = destination_for(ibge, code, index)
        if destination is None:
            validation_only.append(code)
            continue
        row, column, letter = destination
        cell = ws.cell(row=row, column=column)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise RuntimeError(f"Formula protegida em {cell.coordinate}")
        cell.value = round(float(value), 2)
        cell.number_format = '#,##0.00'
        written.append(f"{code}:{letter}{row}")
    return {"written": written, "validation_only": validation_only}
