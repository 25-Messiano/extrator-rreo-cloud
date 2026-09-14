from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader

from database.db import init_db
from services.seed import seed_initial_data
from relatorios.movimento_financeiro import movimento_mes, gerar_pdf_movimento, gerar_excel_movimento


def _data():
    init_db(); seed_initial_data()
    return movimento_mes(2026, 1)


def test_movimento_principal_5_paginas_e_totais():
    d = _data()
    assert len(d['rows']) == 87
    assert str(d['saldo_inicio_geral']) == '13958.55'
    assert str(d['saldo_fim_geral']) == '28416.30'
    assert str(d['totals']['entrada_banco']) == '157286.59'
    assert str(d['totals']['saida_banco']) == '145878.71'
    assert str(d['totals']['entrada_caixa']) == '19859.14'
    assert str(d['totals']['saida_caixa']) == '16809.27'
    pdf = gerar_pdf_movimento(d, 'modelo')
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) == 5
    box = reader.pages[0].mediabox
    assert float(box.width) > float(box.height)


def test_excel_movimento_uma_faixa_e_codigos_com_zero():
    d = _data()
    xlsx = gerar_excel_movimento(d, 'modelo')
    wb = load_workbook(BytesIO(xlsx))
    ws = wb.active
    assert ws.page_setup.orientation == 'landscape'
    assert int(ws.page_setup.paperSize) == 9
    assert ws.page_setup.fitToWidth == 1
    assert len(ws.col_breaks.brk) == 0
    assert any(row[1].value == '0001' for row in ws.iter_rows(min_col=1,max_col=3))
    assert all(c.value != '#REF!' for row in ws.iter_rows() for c in row)
