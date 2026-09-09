from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from integrations.openai_rreo import enabled as openai_enabled, review_rreo_json
from modules.rreo_json.identity import identify_municipality
from modules.rreo_json.extraction import extract_pdf_to_record
from modules.rreo_json.validation import validate_extraction
from modules.rreo_json.destination import build_destination_record, apply_destination_record

ROOT = Path(__file__).resolve().parents[1]
GOIAS = ROOT / 'data' / '20_RREO_FNDE_goias_2025_MODELO_OFICIAL.xlsx'
AGUA = ROOT / 'data' / 'referencias_rreo' / 'RREO_Municipal_2025_Agua_Branca-PB.pdf'


def test_identity_is_name_first_and_ibge_after_match():
    wb = load_workbook(GOIAS, data_only=False)
    try:
        ws = wb['Estado de Goiás-GO']
        result = identify_municipality(
            filename='RREO_Municipal_2025_Abadia de Goias-GO.pdf',
            internal_text='PREFEITURA MUNICIPAL DE ABADIA DE GOIAS\nRELATORIO RESUMIDO',
            ws=ws,
            uf='GO',
        )
        assert result['identificado'] is True
        assert result['melhor_candidato']['ente_planilha'] == 'Abadia de Goiás/GO'
        assert result['melhor_candidato']['ibge'] == '5200050'
        assert result['melhor_candidato']['row'] == 895
        assert 'IBGE SOMENTE APOS HARMONIZACAO' in result['regra']
    finally:
        wb.close()


def test_agua_branca_json_extraction_and_structural_validation():
    extraction = extract_pdf_to_record(AGUA, ano=2025, bimestre=6)
    assert extraction['valores']['1.4'] == 1458429.92
    identity = {'identificado': True, 'melhor_candidato': {'ibge':'2500106','row':1,'ente_planilha':'Água Branca/PB'}}
    validation = validate_extraction(extraction, identity)
    assert validation['status'] == 'VALIDADO_SAFE_JSON'
    assert validation['checks']['2.1.1+2.1.2=2.1']['ok'] is True


def test_destination_is_separate_and_header_driven():
    wb = load_workbook(GOIAS, data_only=False)
    try:
        ws = wb['Estado de Goiás-GO']
        identity = {
            'identificado': True,
            'melhor_candidato': {
                'ibge':'5200050', 'row':895,
                'ente_planilha':'Abadia de Goiás/GO',
            },
        }
        validation = {'status':'VALIDADO_SAFE_JSON'}
        values = {'1.1': 100.0, '1.4': 200.0, '2.1': 999.0, '2.1.1': 500.0, '2.1.2': 499.0, '6.2': 10.0, '6.2.1': 10.0}
        dest = build_destination_record(ws, identity, validation, values)
        assert dest['destinos']['1.1']['celula'] == 'P895'
        assert dest['destinos']['1.4']['celula'] == 'Q895'
        assert dest['destinos']['2.1']['modo'] == 'VALIDACAO_SOMENTE'
        assert dest['destinos']['6.2']['modo'] == 'VALIDACAO_SOMENTE'
        written = apply_destination_record(ws, dest, validation)
        assert 'P895' in written and 'Q895' in written
        assert ws['P895'].value == 100.0
        assert ws['Q895'].value == 200.0
    finally:
        wb.close()


def test_openai_is_optional_and_advisory_without_key(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    assert openai_enabled() is False
    out = review_rreo_json({'valores': {'1.1': 1.0}})
    assert out['advisory_only'] is True
    assert out['status'] == 'SEM_CHAVE'
