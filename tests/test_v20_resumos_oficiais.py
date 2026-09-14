from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader

from database.db import init_db
from services.seed import seed_initial_data
from relatorios.resumos_oficiais import (
    resumo_codigo_ano, suprimento_caixa_ano,
    gerar_pdf_resumo, gerar_excel_resumo,
    gerar_pdf_suprimento, gerar_excel_suprimento,
    gerar_pdf_saldos, gerar_excel_saldos,
)


def _seed():
    init_db(); seed_initial_data()


def test_resumo_banco_e_caixa_conferem_com_base():
    _seed()
    for bc in ('B','C',None):
        d=resumo_codigo_ano(2026,bc)
        assert d['ok']
        assert all(v==0 for v in d['diferencas_entrada'].values())
        assert all(v==0 for v in d['diferencas_saida'].values())
    b=resumo_codigo_ano(2026,'B')
    assert str(b['totais_entrada'][1])=='157286.59'
    assert str(b['totais_saida'][1])=='145878.71'
    c=resumo_codigo_ano(2026,'C')
    assert str(c['totais_entrada'][1])=='19859.14'
    assert str(c['totais_saida'][1])=='16809.27'


def test_pdf_excel_resumos_geram_sem_erro():
    _seed()
    for bc,titulo in [('B','RESUMO GERAL - BANCO'),('C','RESUMO GERAL - CAIXA'),(None,'RESUMO GERAL - BANCO E CAIXA')]:
        pdf=gerar_pdf_resumo(2026,bc,titulo)
        assert len(PdfReader(BytesIO(pdf)).pages)>=2
        wb=load_workbook(BytesIO(gerar_excel_resumo(2026,bc,titulo)))
        assert wb.sheetnames==['1o Semestre','2o Semestre']


def test_suprimento_confronta_banco_caixa_e_exporta():
    _seed();d=suprimento_caixa_ano(2026)
    assert str(d['banco'][1])=='19800.00'
    assert str(d['caixa'][1])=='19800.00'
    assert d['diferenca'][1]==0
    assert len(PdfReader(BytesIO(gerar_pdf_suprimento(2026))).pages)==2
    wb=load_workbook(BytesIO(gerar_excel_suprimento(2026)))
    assert wb['1o Semestre']['B6'].value in ('OK','DIVERGÊNCIA')


def test_saldos_historicos_exportam_tres_quadros():
    _seed()
    pdf=gerar_pdf_saldos(2025,2026)
    assert len(PdfReader(BytesIO(pdf)).pages)>=1
    wb=load_workbook(BytesIO(gerar_excel_saldos(2025,2026)))
    ws=wb.active
    assert any(c.value and 'SALDOS BANCO' in str(c.value) for row in ws.iter_rows() for c in row)
    assert any(c.value and 'SALDOS CAIXA' in str(c.value) for row in ws.iter_rows() for c in row)
