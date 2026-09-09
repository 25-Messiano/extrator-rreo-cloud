from openpyxl import Workbook

from core.tipos_rodada import (
    TipoRodada,
    atualiza_atividade_global,
    forca_processamento,
    limpa_fonte_antes_de_gravar,
    usa_master_existente,
)
from modules.mapeamento_nova_planilha import (
    ABA_DESTINO,
    limpar_campos_fonte,
)


def test_politica_dos_tres_tipos_de_rodada():
    assert not usa_master_existente(TipoRodada.NOVA)
    assert forca_processamento(TipoRodada.NOVA)
    assert not atualiza_atividade_global(TipoRodada.NOVA)

    assert usa_master_existente(TipoRodada.CORRECAO)
    assert forca_processamento(TipoRodada.CORRECAO)
    assert limpa_fonte_antes_de_gravar(TipoRodada.CORRECAO)
    assert atualiza_atividade_global(TipoRodada.CORRECAO)

    assert usa_master_existente(TipoRodada.INCREMENTACAO)
    assert not forca_processamento(TipoRodada.INCREMENTACAO)
    assert not limpa_fonte_antes_de_gravar(TipoRodada.INCREMENTACAO)
    assert atualiza_atividade_global(TipoRodada.INCREMENTACAO)


def test_correcao_limpa_so_a_fonte_escolhida():
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_DESTINO
    row = 2
    # RREO: P/Q; FNDE: V/W
    ws.cell(row=row, column=16).value = 111.0
    ws.cell(row=row, column=17).value = 222.0
    ws.cell(row=row, column=22).value = 333.0
    ws.cell(row=row, column=23).value = 444.0

    limpas = limpar_campos_fonte(ws, row, "FNDE")
    assert set(limpas) == {"V2", "W2"}
    assert ws["P2"].value == 111.0
    assert ws["Q2"].value == 222.0
    assert ws["V2"].value is None
    assert ws["W2"].value is None


def test_rodada_nova_tem_destino_separado_do_master():
    from integrations.google_storage import master_blob_name, round_blob_name
    master = master_blob_name(2025)
    rodada = round_blob_name(2025, "RREO_BRASIL_2025_RODADA_NOVA_20260823_170000.xlsx")
    assert master != rodada
    assert "/RODADAS/2025/" in rodada
