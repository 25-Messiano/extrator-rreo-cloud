from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook, load_workbook

from modules.rreo_json.identity import identify_municipality
from modules.rreo_structure.mapeamento_integrado import load_config, resolve_destination, validate_headers

HEADERS = {
    'A':'Nº','B':'Qut. Municípios','C':'Código IBGE','D':'Ente Federado',
    'E':'2.1.1- Parcela referente à CF, art. 159, I, alínea b',
    'F':'2.1.2- Parcela referente à CF, art. 159, I, alíneas d e e',
    'G':'FPE','H':'2.3- Cota-Parte IPI-Exportação','I':'AFE','J':'2.4- Cota-Parte ITR',
    'K':'2.2- Cota-Parte ICMS','L':'2.5- Cota-Parte IPVA','M':'ITCMD',
    'N':'FUNDEB - Impostos e Transferências de Impostos - 6.1.1 - Principal',
    'O':'FUNDEB - Complementação da União - VAAF - 6.2.1 - Principal',
    'P':'1.1- Receita Resultante do IPTU','Q':'1.4- Receita Resultante do IRRF',
    'R':'1.2- Receita Resultante do ITBI','S':'1.3- Receita Resultante do ISS',
    'T':'2.6- Cota-Parte IOF-Ouro','U':'Petróleo e Gás','V':'PNAE','W':'PNATE','X':'PDDE','Y':'QSE','Z':'PNLD','AA':'Total'
}

def make_book(tmp_path: Path) -> Path:
    p=tmp_path/'modelo.xlsx'; wb=Workbook(); ws=wb.active; ws.title='Estado de Goiás-GO'
    for letter, value in HEADERS.items(): ws[f'{letter}1']=value
    ws['A2']=1; ws['B2']=2; ws['C2']=3; ws['D2']=4
    ws['A3']=233; ws['B3']=25; ws['C3']=2902104; ws['D3']='Araci/BA'
    ws['A4']=999; ws['B4']=200; ws['C4']=2515401; ws['D4']='São Vicente do Seridó/PB'
    wb.save(p); return p

def test_header_map_and_araci_s235_logic(tmp_path):
    p=make_book(tmp_path); wb=load_workbook(p); ws=wb.active
    mp=validate_headers(ws, load_config())
    assert mp['1.3']=='S'; assert mp['2.1'] is None; assert mp['6.2'] is None
    identity={'identificado':True,'status':'EXATO_EXTERNO','melhor_candidato':{'row':3,'ibge':'2902104','ente_planilha':'Araci/BA'}}
    d=resolve_destination(ws, identity, '1.3')
    assert d.celula=='S3'
    wb.close()

def test_external_exact_wins_over_abbreviated_internal(tmp_path):
    p=make_book(tmp_path); wb=load_workbook(p); ws=wb.active
    ident=identify_municipality(filename='RREO_Municipal_2025_São Vicente do Seridó-PB.pdf', internal_text='SERIDÓ - PB\nRREO ANEXO 8', ws=ws, uf='PB')
    assert ident['identificado'] is True
    assert ident['status']=='EXATO_EXTERNO'
    assert ident['melhor_candidato']['ente_planilha']=='São Vicente do Seridó/PB'
    wb.close()

def test_validators_have_no_excel_destination(tmp_path):
    p=make_book(tmp_path); wb=load_workbook(p); ws=wb.active
    identity={'identificado':True,'status':'EXATO_EXTERNO','melhor_candidato':{'row':3,'ibge':'2902104','ente_planilha':'Araci/BA'}}
    assert resolve_destination(ws, identity, '2.1').celula is None
    assert resolve_destination(ws, identity, '6.2').celula is None
    wb.close()

def test_column_a_is_blocked_and_b_authorized():
    cfg=load_config()
    assert cfg['colunas_pdf']['a']['coletar'] is False
    assert cfg['colunas_pdf']['a']['gravar_excel'] is False
    assert cfg['colunas_pdf']['b']['coletar'] is True
    assert cfg['colunas_pdf']['b']['gravar_excel'] is True
