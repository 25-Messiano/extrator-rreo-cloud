from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from modules.rreo import extract_codes, extract_text
from modules.rreo_fields.coluna_previsao_atualizada_a import COLLECT as COLLECT_A, WRITE_TO_EXCEL as WRITE_A
from modules.rreo_fields.coluna_receitas_realizadas_b import COLLECT as COLLECT_B, WRITE_TO_EXCEL as WRITE_B
from modules.rreo_fields.excel_map import CODE_TO_EXCEL, build_ibge_row_index, destination_for
from modules.rreo_fields.registry import MODULE_BY_CODE, extract_many

ROOT = Path(__file__).resolve().parents[1]
AGUA = ROOT / 'data' / 'referencias_rreo' / 'RREO_Municipal_2025_Agua_Branca-PB.pdf'
BASE = ROOT / 'data' / 'RREO-TCM+FNDE PLANILHA BASE.xlsx'


def test_15_codes_have_15_exclusive_modules():
    assert len(MODULE_BY_CODE) == 15
    assert set(MODULE_BY_CODE) == {"1.1","1.2","1.3","1.4","2.1","2.1.1","2.1.2","2.2","2.3","2.4","2.5","2.6","6.1.1","6.2","6.2.1"}


def test_column_a_is_registered_but_never_collectable_or_writable():
    assert COLLECT_A is False
    assert WRITE_A is False
    assert COLLECT_B is True
    assert WRITE_B is True


def test_agua_branca_1_4_cannot_invade_total_2():
    text = extract_text(AGUA)
    values = extract_codes(text, ["1.1","1.2","1.3","1.4","2.1","2.1.1"])
    assert values["1.1"] == 89888.42
    assert values["1.2"] == 45223.58
    assert values["1.3"] == 653260.14
    assert values["1.4"] == 1458429.92
    assert values["1.4"] != 30087736.26
    assert values["2.1"] == 25889148.95
    assert values["2.1.1"] == 22898228.45


def test_exact_block_stops_on_integer_total_code():
    synthetic = '''1.4- Receita Resultante do Imposto de Renda Retido na Fonte - IRRF 974.985,00 1.458.429,92\n2- RECEITA DE TRANSFERENCIAS 24.143.040,00 30.087.736,26\n2.1- Cota-Parte FPM 20.700.000,00 25.889.148,95'''
    values, evidence = extract_many(synthetic, ["1.4", "2.1"])
    assert values["1.4"] == 1458429.92
    assert "30.087.736,26" not in evidence["1.4"]["row_block"]
    assert values["2.1"] == 25889148.95


def test_row_code_prefixes_do_not_collide():
    synthetic = '''2.1- Cota-Parte FPM 20.700.000,00 25.889.148,95\n2.1.1- Parcela referente a CF art 159 alinea b 19.200.000,00 22.898.228,45\n2.1.2- Parcela referente a CF art 159 alineas d e e 1.500.000,00 2.990.920,50'''
    values, _ = extract_many(synthetic, ["2.1", "2.1.1", "2.1.2"])
    assert values == {"2.1": 25889148.95, "2.1.1": 22898228.45, "2.1.2": 2990920.50}


def test_excel_map_matches_official_base_and_does_not_invent_totals():
    assert CODE_TO_EXCEL["2.1"] is None
    assert CODE_TO_EXCEL["6.2"] is None
    wb = load_workbook(BASE, read_only=True, data_only=False)
    try:
        ws = wb['Sequência_PlanilhaCálculo (2)']
        index = build_ibge_row_index(ws)
        # Agua Branca/PB IBGE 2500106
        row = index['2500106']
        assert ws.cell(row=row, column=4).value == 'Água Branca/PB'
        assert destination_for('2500106','1.4',index) == (row,17,'Q')
        assert destination_for('2500106','2.1',index) is None
    finally:
        wb.close()
