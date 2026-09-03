from __future__ import annotations

import io

from openpyxl import load_workbook

from core.central_correcoes import _parse_planilha_estadual, generate_pending_report


def test_parse_planilha_estadual():
    item = {
        "name": "RREO_ESTADO_GO_20260903_120000.xlsx",
        "blob_name": "x/RREO_ESTADO_GO_20260903_120000.xlsx",
        "updated": None,
    }
    parsed = _parse_planilha_estadual(item)
    assert parsed is not None
    assert parsed.uf == "GO"
    assert parsed.fonte == "RREO"


def test_parse_ignora_master_e_multiestado():
    assert _parse_planilha_estadual({"name": "RREO_FNDE_BRASIL_MASTER_2025.xlsx"}) is None
    assert _parse_planilha_estadual({"name": "RREO_SELECIONADOS_RR-AP_2025.xlsx"}) is None


def test_relatorio_pendencias(tmp_path):
    out = generate_pending_report(2025, [{
        "uf": "GO", "codigo_ibge": "5215702", "municipio": "Palmeiras de Goiás",
        "status_geral": "PENDENTE", "status_rreo": "PENDENTE", "pdf_rreo": True,
        "pdf_rreo_nome": "RREO.pdf", "campos_rreo": 0, "esperados_rreo": 13,
        "status_fnde": "NA", "pdf_fnde": None, "pdf_fnde_nome": "",
        "campos_fnde": 0, "esperados_fnde": 4, "planilhas_origem": [],
    }], tmp_path / "p.xlsx")
    wb = load_workbook(out, read_only=True, data_only=True)
    ws = wb["PENDENCIAS"]
    assert ws.cell(2, 4).value == "Palmeiras de Goiás"
    assert ws.cell(2, 7).value == "SIM"
    wb.close()

from openpyxl import Workbook
from core.central_correcoes import _read_state_sheet, PlanilhaEstadual
from modules.mapeamento_nova_planilha import ABA_DESTINO
from core.auditoria_rreo import RREO_LOG_HEADERS, RREO_LOG_SHEET


def test_reconstrucao_preserva_erro_do_log():
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_DESTINO
    ws.cell(1, 3, "Código IBGE")
    ws.cell(1, 4, "Ente Federado")
    ws.cell(3, 3, 5215702)
    ws.cell(3, 4, "Palmeiras de Goiás/GO")
    log = wb.create_sheet(RREO_LOG_SHEET)
    log.append(RREO_LOG_HEADERS)
    values = {header: "" for header in RREO_LOG_HEADERS}
    values["Código IBGE"] = "5215702"
    values["Status"] = "ERRO"
    values["Erro resumido"] = "PDF divergente"
    log.append([values[h] for h in RREO_LOG_HEADERS])
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    origin = PlanilhaEstadual("GO", "RREO", "RREO_ESTADO_GO.xlsx", "x", None)
    rows = _read_state_sheet(bio.getvalue(), "GO", "RREO", origin)
    assert rows[0]["status_rreo"] == "ERRO"
    assert rows[0]["erro_rreo"] == "PDF divergente"
