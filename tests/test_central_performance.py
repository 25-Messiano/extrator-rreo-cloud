from __future__ import annotations

import io
import time
from datetime import datetime, timezone

from openpyxl import Workbook

from core.central_correcoes import PlanilhaEstadual, _read_state_sheet
from modules.mapeamento_nova_planilha import ABA_DESTINO


def _xlsx(rows: int, uf: str = "MG") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_DESTINO
    ws.cell(1, 1).value = "TESTE"
    # starts row 3; col C IBGE, col D Ente/UF
    for i in range(rows):
        r = i + 3
        ws.cell(r, 3).value = f"31{(i+1):05d}"[:7]
        ws.cell(r, 4).value = f"Municipio {i+1}/{uf}"
        for c in range(5, 20):
            ws.cell(r, c).value = float(i + c)
    bio = io.BytesIO(); wb.save(bio); wb.close(); return bio.getvalue()


def test_leitura_streaming_853_registros_em_tempo_razoavel():
    payload = _xlsx(853)
    origin = PlanilhaEstadual("MG", "RREO", "RREO_MG.xlsx", "x/RREO_MG.xlsx", datetime.now(timezone.utc))
    started = time.perf_counter()
    rows = _read_state_sheet(payload, "MG", "RREO", origin)
    elapsed = time.perf_counter() - started
    assert len(rows) == 853
    assert elapsed < 8.0
