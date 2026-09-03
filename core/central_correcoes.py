from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from core.identificacao_arquivos import identificar_uf
from core.auditoria_rreo import RREO_LOG_SHEET
from core.auditoria_fnde import FNDE_LOG_SHEET
from core.indice_rreo import build_rreo_index, localizar_por_municipio
from integrations.google_storage import (
    RESULTADOS_PREFIX,
    download_bytes,
    find_fnde_folder,
    find_rreo_folder,
    list_fnde_pdfs_by_uf,
    list_results,
    list_rreo_pdfs_by_uf,
    upload_file,
    upload_text,
)
from modules.mapeamento_nova_planilha import ABA_DESTINO, CAMPOS_DESTINO

CENTRAL_PREFIX = f"{RESULTADOS_PREFIX}CENTRAL_CORRECOES/"
CATALOGO_SCHEMA = "CENTRAL_CORRECOES_V1"

_RREO_COLS = [campo.coluna for campo in CAMPOS_DESTINO if campo.fonte == "RREO"]
_FNDE_COLS = [campo.coluna for campo in CAMPOS_DESTINO if campo.fonte == "FNDE"]


@dataclass(frozen=True)
class PlanilhaEstadual:
    uf: str
    fonte: str
    name: str
    blob_name: str
    updated: Any


def _agora() -> str:
    return datetime.now(timezone.utc).isoformat()


def _catalog_blob(year: int) -> str:
    return f"{CENTRAL_PREFIX}{int(year)}/catalogo.json"


def _state_blob(year: int, uf: str) -> str:
    return f"{CENTRAL_PREFIX}{int(year)}/JSON/{str(uf).upper()}.json"


def _report_blob(year: int, filename: str) -> str:
    return f"{CENTRAL_PREFIX}{int(year)}/RELATORIOS/{Path(filename).name}"


def _parse_planilha_estadual(item: dict[str, Any]) -> PlanilhaEstadual | None:
    """Reconhece somente saídas estaduais independentes.

    Master, rodadas multiestado e relatórios ficam fora para evitar misturar
    resultados de naturezas diferentes ao reconstruir o índice retroativo.
    """
    name = str(item.get("name") or "")
    upper = name.upper()
    if not upper.endswith(".XLSX"):
        return None
    if "MASTER" in upper or "TODOS_OS_ESTADOS" in upper or "SELECIONADOS" in upper:
        return None

    source = ""
    if upper.startswith("RREO_FNDE_"):
        source = "RREO+FNDE"
    elif upper.startswith("RREO_"):
        source = "RREO"
    elif upper.startswith("FNDE_"):
        source = "FNDE"
    else:
        return None

    # Há duas famílias reais de saídas estaduais no app:
    # 1) RREO_ESTADO_GO_...xlsx (fluxos antigos/checkpoints)
    # 2) RREO_GO_2025_B6_RODADA_NOVA_...xlsx (Rodada Nova atual)
    # A v1.3.2 aceitava apenas a primeira e por isso podia reconstruir 0 UFs.
    match = re.search(r"_ESTADO_([A-Z]{2})(?:_|\.)", upper)
    if not match and "_RODADA_NOVA_" in upper:
        match = re.match(r"^(?:RREO_FNDE|RREO|FNDE)_([A-Z]{2})_\d{4}(?:_|\.)", upper)
    if match:
        uf = match.group(1)
    else:
        # Checkpoints podem estar armazenados em .../03_PLANILHAS_PROCESSADAS/UF/.
        blob_name = str(item.get("blob_name") or "").upper()
        folder_match = re.search(r"/03_PLANILHAS_PROCESSADAS/([A-Z]{2})/", blob_name)
        uf = folder_match.group(1) if folder_match else identificar_uf(name)
    if not uf:
        return None
    return PlanilhaEstadual(
        uf=uf,
        fonte=source,
        name=name,
        blob_name=str(item.get("blob_name") or ""),
        updated=item.get("updated"),
    )


def _discover_master(year: int) -> PlanilhaEstadual | None:
    """Localiza o master cumulativo do ano como fallback de reconstrução.

    Muitos estados foram processados individualmente em modo Incrementação/Correção,
    que atualiza o MASTER em vez de produzir um arquivo estadual independente.
    """
    expected = f"RREO_FNDE_BRASIL_MASTER_{int(year)}.XLSX"
    candidates: list[PlanilhaEstadual] = []
    for item in list_results(None):
        name = str(item.get("name") or "")
        if name.upper() != expected:
            continue
        candidates.append(PlanilhaEstadual(
            uf="BR", fonte="RREO+FNDE", name=name,
            blob_name=str(item.get("blob_name") or ""), updated=item.get("updated"),
        ))
    if not candidates:
        return None
    floor = datetime.min.replace(tzinfo=timezone.utc)
    return max(candidates, key=lambda x: x.updated or floor)


def _ufs_from_workbook(payload: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        if ABA_DESTINO not in wb.sheetnames:
            return []
        ws = wb[ABA_DESTINO]
        ufs: set[str] = set()
        for row in range(3, ws.max_row + 1):
            ente = str(ws.cell(row=row, column=4).value or "").strip().upper()
            match = re.search(r"/([A-Z]{2})$", ente)
            if match:
                ufs.add(match.group(1))
        return sorted(ufs)
    finally:
        wb.close()


def discover_latest_state_spreadsheets(year: int) -> dict[tuple[str, str], PlanilhaEstadual]:
    """Obtém a planilha estadual mais recente por UF e fonte."""
    del year  # nomes atuais não carregam o ano de forma confiável; o conteúdo é a autoridade.
    latest: dict[tuple[str, str], PlanilhaEstadual] = {}
    for item in list_results(None):
        parsed = _parse_planilha_estadual(item)
        if parsed is None:
            continue
        key = (parsed.uf, parsed.fonte)
        previous = latest.get(key)
        if previous is None or (parsed.updated or datetime.min.replace(tzinfo=timezone.utc)) > (
            previous.updated or datetime.min.replace(tzinfo=timezone.utc)
        ):
            latest[key] = parsed
    return latest


def _value_present(value: Any) -> bool:
    return value is not None and value != ""


def _source_status(ws, row: int, columns: list[int]) -> tuple[str, int, int]:
    present = sum(1 for col in columns if _value_present(ws.cell(row=row, column=col).value))
    expected = len(columns)
    if present == 0:
        return "PENDENTE", present, expected
    if present < expected:
        return "PARCIAL", present, expected
    return "DADOS_PRESENTES", present, expected


def _latest_log_by_ibge(wb, sheet_name: str) -> dict[str, dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return {}
    ws = wb[sheet_name]
    headers = {str(cell.value or "").strip(): idx for idx, cell in enumerate(ws[1], start=1)}
    code_col = headers.get("Código IBGE")
    if not code_col:
        return {}
    out: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        code = str(ws.cell(row=row, column=code_col).value or "").strip().split(".")[0]
        if len(code) != 7 or not code.isdigit():
            continue
        item = {header: ws.cell(row=row, column=col).value for header, col in headers.items()}
        out[code] = item
    return out


def _apply_log_status(item: dict[str, Any], log: dict[str, Any] | None, source: str) -> None:
    if not log:
        return
    raw_status = str(log.get("Status") or "").upper().strip()
    if raw_status in {"OK", "ERRO", "PARCIAL"}:
        item[f"status_{source.lower()}"] = raw_status
    item[f"erro_{source.lower()}"] = str(log.get("Erro resumido") or "").strip()
    item[f"pdf_{source.lower()}_log"] = str(log.get(f"PDF {source}") or "").strip()
    item[f"metodo_{source.lower()}"] = str(log.get("Método de extração") or "").strip()
    item[f"validacao_{source.lower()}"] = str(log.get("Validação dupla") or "").strip()


def _read_state_sheet(payload: bytes, uf: str, source: str, origin: PlanilhaEstadual) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        if ABA_DESTINO not in wb.sheetnames:
            raise ValueError(f"Aba obrigatória ausente: {ABA_DESTINO}")
        ws = wb[ABA_DESTINO]
        rreo_logs = _latest_log_by_ibge(wb, RREO_LOG_SHEET)
        fnde_logs = _latest_log_by_ibge(wb, FNDE_LOG_SHEET)
        records: list[dict[str, Any]] = []
        for row in range(3, ws.max_row + 1):
            code_raw = ws.cell(row=row, column=3).value
            ente = str(ws.cell(row=row, column=4).value or "").strip()
            code = str(code_raw or "").strip().split(".")[0]
            if len(code) != 7 or not code.isdigit() or not ente.upper().endswith(f"/{uf}"):
                continue
            municipio = ente.rsplit("/", 1)[0].strip()
            item = {
                "ano": None,
                "codigo_ibge": code,
                "uf": uf,
                "municipio": municipio,
                "row": row,
                "planilha_origem": origin.name,
                "blob_origem": origin.blob_name,
                "fonte_planilha": source,
                "status_rreo": "NA",
                "status_fnde": "NA",
                "campos_rreo": 0,
                "campos_fnde": 0,
                "esperados_rreo": len(_RREO_COLS),
                "esperados_fnde": len(_FNDE_COLS),
            }
            if source in {"RREO", "RREO+FNDE"}:
                status, count, expected = _source_status(ws, row, _RREO_COLS)
                item.update(status_rreo=status, campos_rreo=count, esperados_rreo=expected)
            if source in {"FNDE", "RREO+FNDE"}:
                status, count, expected = _source_status(ws, row, _FNDE_COLS)
                item.update(status_fnde=status, campos_fnde=count, esperados_fnde=expected)
            if source in {"RREO", "RREO+FNDE"}:
                _apply_log_status(item, rreo_logs.get(code), "RREO")
            if source in {"FNDE", "RREO+FNDE"}:
                _apply_log_status(item, fnde_logs.get(code), "FNDE")
            records.append(item)
        return records
    finally:
        wb.close()


def _merge_record(target: dict[str, Any], incoming: dict[str, Any]) -> None:
    for source in ("rreo", "fnde"):
        status_key = f"status_{source}"
        count_key = f"campos_{source}"
        if incoming.get(status_key) != "NA":
            target[status_key] = incoming.get(status_key)
            target[count_key] = incoming.get(count_key, 0)
            for suffix in ("erro", "pdf", "metodo", "validacao"):
                key = f"{suffix}_{source}" if suffix != "pdf" else f"pdf_{source}_log"
                if incoming.get(key):
                    target[key] = incoming.get(key)
    origins = target.setdefault("planilhas_origem", [])
    origin = {
        "name": incoming.get("planilha_origem", ""),
        "blob_name": incoming.get("blob_origem", ""),
        "fonte": incoming.get("fonte_planilha", ""),
    }
    if origin not in origins:
        origins.append(origin)


def _overall_status(record: dict[str, Any]) -> str:
    statuses = [record.get("status_rreo", "NA"), record.get("status_fnde", "NA")]
    active = [s for s in statuses if s != "NA"]
    if not active:
        return "SEM_DADOS"
    if any(s == "ERRO" for s in active):
        return "ERRO"
    if any(s == "PENDENTE" for s in active):
        return "PENDENTE"
    if any(s == "PARCIAL" for s in active):
        return "PARCIAL"
    return "DADOS_PRESENTES"


def _annotate_pdf_availability(records: list[dict[str, Any]], uf: str, year: int, bimestre: int = 6) -> None:
    """Confere existência dos PDFs pelos índices externos, sem abrir PDFs."""
    by_code = {str(item["codigo_ibge"]): item for item in records}
    municipios = [
        {"codigo_ibge": r["codigo_ibge"], "nome": r["municipio"], "uf": uf, "row": r.get("row", 0)}
        for r in records
    ]
    try:
        files = list_rreo_pdfs_by_uf(uf, year=year, bimestre=bimestre)
        index = build_rreo_index(files, uf=uf, municipios=municipios)
        for code, record in by_code.items():
            found = localizar_por_municipio(index, record["municipio"], uf, code)
            record["pdf_rreo"] = bool(found)
            record["pdf_rreo_nome"] = str((found or {}).get("name") or "")
    except Exception as error:
        for record in records:
            record["pdf_rreo"] = None
            record["pdf_rreo_erro"] = str(error)

    try:
        fnde_files = list_fnde_pdfs_by_uf(uf, year=year)
        fnde_by_code: dict[str, dict[str, Any]] = {}
        for file in fnde_files:
            text = str(file.get("name") or "")
            match = re.search(r"(?<!\d)(\d{7})(?!\d)", text)
            if match:
                fnde_by_code.setdefault(match.group(1), file)
        for code, record in by_code.items():
            found = fnde_by_code.get(code)
            record["pdf_fnde"] = bool(found)
            record["pdf_fnde_nome"] = str((found or {}).get("name") or "")
    except Exception as error:
        for record in records:
            record["pdf_fnde"] = None
            record["pdf_fnde_erro"] = str(error)


def rebuild_index_from_state_spreadsheets(
    year: int,
    *,
    progress: Callable[[int, int, str], None] | None = None,
    annotate_pdfs: bool = True,
    bimestre: int = 6,
) -> dict[str, Any]:
    """Reconstrói JSONs por estado a partir das planilhas estaduais existentes.

    Não altera nenhuma planilha processada. A operação é somente leitura das
    planilhas e gravação de um novo catálogo técnico em CENTRAL_CORRECOES.
    """
    latest = discover_latest_state_spreadsheets(year)
    master = _discover_master(year)
    master_payload: bytes | None = None
    master_ufs: list[str] = []
    if master is not None:
        try:
            master_payload = download_bytes(master.blob_name)
            master_ufs = _ufs_from_workbook(master_payload)
        except Exception:
            master_payload = None
            master_ufs = []
    ufs = sorted({uf for uf, _ in latest} | set(master_ufs))
    merged_all: list[dict[str, Any]] = []
    total = len(ufs)
    for pos, uf in enumerate(ufs, start=1):
        if progress:
            progress(pos, total, uf)
        merged: dict[str, dict[str, Any]] = {}
        sources_used: list[dict[str, str]] = []
        for source in ("RREO", "FNDE", "RREO+FNDE"):
            plan = latest.get((uf, source))
            if plan is None:
                continue
            payload = download_bytes(plan.blob_name)
            records = _read_state_sheet(payload, uf, source, plan)
            sources_used.append({"fonte": source, "name": plan.name, "blob_name": plan.blob_name})
            for item in records:
                code = item["codigo_ibge"]
                if code not in merged:
                    merged[code] = {
                        "ano": int(year),
                        "codigo_ibge": code,
                        "uf": uf,
                        "municipio": item["municipio"],
                        "row": item.get("row", 0),
                        "status_rreo": "NA",
                        "status_fnde": "NA",
                        "campos_rreo": 0,
                        "campos_fnde": 0,
                        "esperados_rreo": len(_RREO_COLS),
                        "esperados_fnde": len(_FNDE_COLS),
                        "planilhas_origem": [],
                    }
                _merge_record(merged[code], item)
        # Se não houver planilha estadual reconhecida para a UF, usa o MASTER
        # cumulativo como fonte retroativa. Isso reaproveita estados processados
        # individualmente em Incrementação/Correção sem reprocessar os PDFs.
        if not merged and master is not None and master_payload is not None and uf in master_ufs:
            records_master = _read_state_sheet(master_payload, uf, "RREO+FNDE", master)
            if records_master:
                sources_used.append({
                    "fonte": "MASTER", "name": master.name, "blob_name": master.blob_name,
                })
                for item in records_master:
                    code = item["codigo_ibge"]
                    merged[code] = {
                        "ano": int(year),
                        "codigo_ibge": code,
                        "uf": uf,
                        "municipio": item["municipio"],
                        "row": item.get("row", 0),
                        "status_rreo": "NA",
                        "status_fnde": "NA",
                        "campos_rreo": 0,
                        "campos_fnde": 0,
                        "esperados_rreo": len(_RREO_COLS),
                        "esperados_fnde": len(_FNDE_COLS),
                        "planilhas_origem": [],
                    }
                    _merge_record(merged[code], item)

        records = list(merged.values())
        if annotate_pdfs:
            _annotate_pdf_availability(records, uf, int(year), int(bimestre))
        for record in records:
            record["status_geral"] = _overall_status(record)
        state_payload = {
            "schema": CATALOGO_SCHEMA,
            "ano": int(year),
            "uf": uf,
            "gerado_em": _agora(),
            "bimestre_rreo": int(bimestre),
            "planilhas_usadas": sources_used,
            "municipios": records,
        }
        upload_text(json.dumps(state_payload, ensure_ascii=False, indent=2), _state_blob(year, uf), "application/json; charset=utf-8")
        merged_all.extend(records)

    catalog = {
        "schema": CATALOGO_SCHEMA,
        "ano": int(year),
        "gerado_em": _agora(),
        "bimestre_rreo": int(bimestre),
        "ufs": ufs,
        "total_municipios": len(merged_all),
        "municipios": merged_all,
    }
    upload_text(json.dumps(catalog, ensure_ascii=False, indent=2), _catalog_blob(year), "application/json; charset=utf-8")
    return catalog


def load_catalog(year: int) -> dict[str, Any] | None:
    try:
        raw = download_bytes(_catalog_blob(year))
        data = json.loads(raw.decode("utf-8"))
        return data if isinstance(data, dict) else None
    except Exception:
        return None


def generate_pending_report(year: int, records: list[dict[str, Any]], destination: str | Path) -> Path:
    path = Path(destination)
    wb = Workbook()
    ws = wb.active
    ws.title = "PENDENCIAS"
    headers = [
        "Ano", "UF", "Código IBGE", "Município", "Status Geral",
        "Status RREO", "PDF RREO", "Arquivo RREO", "Campos RREO", "Erro RREO",
        "Status FNDE", "PDF FNDE", "Arquivo FNDE", "Campos FNDE", "Erro FNDE",
        "Planilhas de origem",
    ]
    ws.append(headers)
    for record in records:
        if record.get("status_geral") == "DADOS_PRESENTES":
            continue
        origins = "; ".join(item.get("name", "") for item in record.get("planilhas_origem", []))
        ws.append([
            int(year), record.get("uf", ""), record.get("codigo_ibge", ""), record.get("municipio", ""),
            record.get("status_geral", ""), record.get("status_rreo", ""),
            "SIM" if record.get("pdf_rreo") is True else ("NÃO" if record.get("pdf_rreo") is False else "N/D"),
            record.get("pdf_rreo_nome", ""), f"{record.get('campos_rreo', 0)}/{record.get('esperados_rreo', 0)}", record.get("erro_rreo", ""),
            record.get("status_fnde", ""),
            "SIM" if record.get("pdf_fnde") is True else ("NÃO" if record.get("pdf_fnde") is False else "N/D"),
            record.get("pdf_fnde_nome", ""), f"{record.get('campos_fnde', 0)}/{record.get('esperados_fnde', 0)}", record.get("erro_fnde", ""),
            origins,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A": 9, "B": 7, "C": 15, "D": 32, "E": 18, "F": 16, "G": 12, "H": 45, "I": 15, "J": 55,
              "K": 16, "L": 12, "M": 45, "N": 15, "O": 55, "P": 60}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wb.save(path)
    wb.close()
    return path


def upload_pending_report(year: int, path: str | Path) -> dict[str, Any]:
    return upload_file(
        path,
        _report_blob(year, Path(path).name),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def correction_preset(year: int, record: dict[str, Any], source: str) -> dict[str, Any]:
    uf = str(record.get("uf") or "").upper()
    source = str(source or "RREO").upper()
    if source == "FNDE":
        state_folder = find_fnde_folder(uf, year) or uf
        execution = "FNDE — Município único"
    elif source in {"RREO+FNDE", "RREO E FNDE"}:
        state_folder = find_rreo_folder(uf, year) or find_fnde_folder(uf, year) or uf
        execution = "RREO + FNDE — Município único"
    else:
        state_folder = find_rreo_folder(uf, year) or uf
        execution = "RREO — Município único"
    return {
        "ano": int(year),
        "uf": uf,
        "codigo_ibge": str(record.get("codigo_ibge") or ""),
        "municipio": str(record.get("municipio") or ""),
        "fonte": source,
        "estado_cloud": state_folder,
        "execucao": execution,
        "tipo_rodada": "Rodada de Correção",
    }


__all__ = [
    "CENTRAL_PREFIX",
    "discover_latest_state_spreadsheets",
    "rebuild_index_from_state_spreadsheets",
    "load_catalog",
    "generate_pending_report",
    "upload_pending_report",
    "correction_preset",
]
